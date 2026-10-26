"""Recalcula em Python, de forma independente, os números do Resumo do painel (v2, parâmetros padrão).

Uso: python tests/conferir_totais.py PASTA [ULTIMO_MES=AAAA-MM] [modo=pag|ref] [unidade=mat|evento]

Regras (as mesmas do painel):
- Termos: arquivo com "termo" no nome, só a aba Termos. Mês do termo = Data de negociação (sem ela, Data de aplicação).
  Termos da mesma matrícula no mesmo mês são um evento só. Unidade "mat" usa o primeiro termo do ano de cada matrícula.
- Mês pago: soma do Valor Pago no mês > 0 (pela data do pagamento ou pela referência Mês/Ano).
- Mês sem dados: mês sem nenhum pagamento na base inteira; fica fora das janelas.
- Antes: 12 meses anteriores ao mês do termo. 10 a 12 pagos (proporcional, ≥ 10/12) Adimplente, ≥ 6/12 Flutuante, > 0 Inadimplente, 0 Nunca pagou.
- Depois: do mês seguinte ao termo até o último mês completo, no máximo 12 meses, e até o mês anterior ao próximo termo
  da mesma matrícula. Menos de 3 meses: acompanhamento insuficiente. ≥ 83,3% Adimplente, ≥ 50% Flutuante, > 0 Inadimplente, 0 Nunca pagou.
"""
import csv
import os
import re
import sys
import unicodedata
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

pasta = sys.argv[1]
ULT = sys.argv[2] if len(sys.argv) > 2 else None
MODO = sys.argv[3] if len(sys.argv) > 3 else 'pag'
UNIDADE = sys.argv[4] if len(sys.argv) > 4 else 'mat'
MIN, ADIM, FLUT = 3, 83.3, 50.0
FAIXAS = ['Adimplente', 'Flutuante', 'Inadimplente', 'Nunca pagou', 'Acomp. insuf.']


def norm(s):
    s = unicodedata.normalize('NFD', str(s or '')).encode('ascii', 'ignore').decode().upper()
    return re.sub(r'[^A-Z0-9]+', '_', s).strip('_')


def lig(v):
    if v is None or v == '':
        return ''
    if isinstance(v, (int, float)):
        return str(round(v))
    s = str(v).strip()
    return str(int(float(s.replace(',', '.')))) if re.fullmatch(r'\d+([.,]0+)?', s) else s


def data(v):
    if v in (None, ''):
        return None
    if isinstance(v, (datetime, date)):
        return v.year * 12 + v.month - 1, v.strftime('%Y%m%d')
    if isinstance(v, (int, float)):
        d = datetime(1899, 12, 30) + timedelta(days=v)
        return d.year * 12 + d.month - 1, d.strftime('%Y%m%d')
    m = re.match(r'(\d{1,2})/(\d{1,2})/(\d{4})', str(v).strip())
    return (int(m[3]) * 12 + int(m[2]) - 1, m[3] + m[2].zfill(2) + m[1].zfill(2)) if m else None


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or '').strip()
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return 0.0


def tabelas(caminho, so_termos=False):
    if caminho.lower().endswith('.csv'):
        raw = open(caminho, 'rb').read()
        try:
            txt = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            txt = raw.decode('cp1252')
        yield list(csv.reader(txt.splitlines(), delimiter=';'))
    else:
        wb = load_workbook(caminho, read_only=True)
        abas = [ws for ws in wb.worksheets if so_termos and norm(ws.title) == 'TERMOS'] or wb.worksheets
        for ws in abas:
            yield list(ws.iter_rows(values_only=True))


def com_cabecalho(linhas):
    for i, l in enumerate(linhas[:30]):
        nomes = [norm(c) for c in l]
        if 'N_LIGACAO' in nomes or 'MATRICULA' in nomes:
            return nomes, linhas[i + 1:]
    return None, []


arquivos = sorted(os.path.join(r, f) for r, _, fs in os.walk(pasta) for f in fs if f.lower().endswith(('.xlsx', '.csv')))
termos = {}   # matrícula -> {mês: data mais antiga}
for a in arquivos:
    if 'termo' not in os.path.basename(a).lower():
        continue
    for t in tabelas(a, so_termos=True):
        h, linhas = com_cabecalho(t)
        if not h:
            continue
        il = h.index('MATRICULA') if 'MATRICULA' in h else h.index('N_LIGACAO')
        ineg, iapl = h.index('DATA_DE_NEGOCIACAO'), h.index('DATA_DE_APLICACAO')
        for l in linhas:
            k, d = lig(l[il]), data(l[ineg]) or data(l[iapl])
            if not k or not d or d[0] // 12 != 2026:
                continue
            ev = termos.setdefault(k, {})
            ev[d[0]] = min(ev.get(d[0], d[1]), d[1])

pag, vistos, min_am, max_am, cobertura = {}, set(), 10 ** 9, -1, set()
for a in arquivos:
    if 'termo' in os.path.basename(a).lower():
        continue
    for t in tabelas(a):
        h, linhas = com_cabecalho(t)
        if not h:
            continue
        ix = {n: h.index(n) for n in ('N_LIGACAO', 'N_NOTA', 'MES', 'ANO', 'VALOR_PAGO', 'DATA_PAGTO')}
        for l in linhas:
            k = lig(l[ix['N_LIGACAO']])
            if not k:
                continue
            ref = int(num(l[ix['ANO']])) * 12 + int(num(l[ix['MES']])) - 1
            dp = data(l[ix['DATA_PAGTO']])
            min_am, max_am = min(min_am, dp[0]), max(max_am, dp[0])
            cobertura.add(dp[0] if MODO == 'pag' else ref)
            if k not in termos:
                continue
            v = num(l[ix['VALOR_PAGO']])
            chave = (k, str(l[ix['N_NOTA']]).strip(), ref, dp[1], v)
            if chave in vistos:
                continue
            vistos.add(chave)
            if v > 0:
                mes = ref if MODO == 'ref' else dp[0]
                pag[(k, mes)] = pag.get((k, mes), 0) + v

ult = int(ULT[:4]) * 12 + int(ULT[5:7]) - 1 if ULT else max_am - 1
ok = lambda m: min_am <= m <= ult and m in cobertura


def antes(p, n):
    if n == 0:
        return -1
    if p == 0:
        return 3
    return 0 if p / n >= 10 / 12 - 1e-9 else 1 if p / n >= .5 - 1e-9 else 2


def depois(p, n):
    if n < MIN:
        return 4
    if p == 0:
        return 3
    x = p * 100 / n
    return 0 if x >= ADIM - 1e-9 else 1 if x >= FLUT - 1e-9 else 2


M = [[0] * 5 for _ in range(4)]
c = dict(total=0, comp=0, melhorou=0, piorou=0, igual=0, insuf=0, semhist=0, naopag=0, recup=0, mA=0.0, mD=0.0)
for k, evs in termos.items():
    meses = sorted(evs)
    for i, t in enumerate(meses):
        if UNIDADE == 'mat' and i > 0:
            break
        c['total'] += 1
        fim = min(ult, t + 12)
        if i + 1 < len(meses):
            fim = min(fim, meses[i + 1] - 1)
        ja = [m for m in range(t - 12, t) if ok(m)]
        jd = [m for m in range(t + 1, fim + 1) if ok(m)]
        va, vd = [pag.get((k, m), 0) for m in ja], [pag.get((k, m), 0) for m in jd]
        ca, cd = antes(sum(v > 0 for v in va), len(va)), depois(sum(v > 0 for v in vd), len(vd))
        if ca < 0:
            c['semhist'] += 1
            continue
        M[ca][cd] += 1
        if cd == 4:
            c['insuf'] += 1
            continue
        c['comp'] += 1
        c['melhorou' if cd < ca else 'piorou' if cd > ca else 'igual'] += 1
        c['mA'] += sum(va) / len(va); c['mD'] += sum(vd) / len(vd)
        if ca >= 2:
            c['naopag'] += 1; c['recup'] += cd <= 1

rot = lambda a: ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez'][a % 12] + '/' + str(a // 12)[2:]
sem = [rot(m) for m in range(min_am, ult + 1) if m not in cobertura]
print(f'Pagamentos de {rot(min_am)} a {rot(max_am)}, análise até {rot(ult)}, modo {MODO}, unidade {UNIDADE}' + (f', meses sem dados: {", ".join(sem)}' if sem else ''))
print(f"Total: {c['total']} | suficiente: {c['comp']} | melhoraram: {c['melhorou']} | pioraram: {c['piorou']} | permaneceram: {c['igual']} | insuficiente: {c['insuf']} | sem histórico: {c['semhist']}")
print(f"Recuperados: {c['recup']} de {c['naopag']} | média mensal por cliente comparável: antes R$ {c['mA'] / max(1, c['comp']):,.2f}, depois R$ {c['mD'] / max(1, c['comp']):,.2f}")
print('Matriz (linhas = antes, colunas = depois):', ' '.join(f'{f[:6]:>7}' for f in FAIXAS))
for a in range(4):
    print(f'  {FAIXAS[a]:<13}', ' '.join(f'{x:>7}' for x in M[a]), f'| {sum(M[a]):>6}')
